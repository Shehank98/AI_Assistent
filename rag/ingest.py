"""
rag/ingest.py — LangChain-powered document ingestion for Jarvis.

Supported formats:
  .xlsx, .xls, .xlsm  — pandas (openpyxl / xlrd engines, handles both old and new Excel)
  .csv                 — pandas
  .pdf                 — LangChain PyPDFLoader (pypdf)
  .docx                — LangChain Docx2txtLoader (docx2txt)
  .pptx                — python-pptx (slide-by-slide extraction)
  .txt, .md, .rst      — LangChain TextLoader
  .html, .htm          — LangChain BSHTMLLoader
  .json                — LangChain JSONLoader
  anything else        — TextLoader fallback
"""

from pathlib import Path


# ── Excel (.xlsx and .xls) ────────────────────────────────────────────────────

def _load_excel(file_path: str) -> list[dict]:
    import pandas as pd

    source = Path(file_path).name
    chunks = []
    ROWS_PER_CHUNK = 25

    # pandas auto-selects engine: openpyxl for .xlsx, xlrd for .xls
    xl = pd.ExcelFile(file_path)
    for sheet_name in xl.sheet_names:
        df = xl.parse(sheet_name)
        if df.empty:
            continue

        # Build row strings
        row_strs = []
        for idx, row in df.iterrows():
            parts = [
                f"{col}: {val}"
                for col, val in row.items()
                if val is not None and str(val).strip() and str(val) != "nan"
            ]
            if parts:
                row_strs.append(f"Row {idx + 1}: " + " | ".join(parts))

        for start in range(0, len(row_strs), ROWS_PER_CHUNK):
            batch = row_strs[start : start + ROWS_PER_CHUNK]
            if batch:
                chunks.append({
                    "content": "\n".join(batch),
                    "source": source,
                    "metadata": {
                        "file": source,
                        "sheet": sheet_name,
                        "row_start": start + 1,
                        "row_end": start + len(batch),
                    },
                })
    return chunks


# ── CSV ───────────────────────────────────────────────────────────────────────

def _load_csv(file_path: str) -> list[dict]:
    import pandas as pd

    source = Path(file_path).name
    chunks = []
    ROWS_PER_CHUNK = 30

    df = pd.read_csv(file_path)
    for start in range(0, len(df), ROWS_PER_CHUNK):
        batch = df.iloc[start : start + ROWS_PER_CHUNK]
        rows = []
        for idx, row in batch.iterrows():
            parts = [
                f"{col}: {val}"
                for col, val in row.items()
                if val is not None and str(val).strip() and str(val) != "nan"
            ]
            if parts:
                rows.append(f"Row {idx + 1}: " + " | ".join(parts))
        if rows:
            chunks.append({
                "content": "\n".join(rows),
                "source": source,
                "metadata": {"file": source, "row_start": start + 1, "row_end": start + len(batch)},
            })
    return chunks


# ── PowerPoint ────────────────────────────────────────────────────────────────

def _load_pptx(file_path: str) -> list[dict]:
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("pip install python-pptx")

    source = Path(file_path).name
    prs = Presentation(file_path)
    chunks = []
    for slide_num, slide in enumerate(prs.slides, 1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if texts:
            chunks.append({
                "content": "\n".join(texts),
                "source": source,
                "metadata": {"file": source, "slide": slide_num},
            })
    return chunks


# ── LangChain-based loaders (PDF, Word, HTML, JSON, plain text) ───────────────

def _lc_load(file_path: str) -> list[dict]:
    """Use a LangChain loader, split with RecursiveCharacterTextSplitter."""
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter
    except ImportError:
        from langchain.text_splitter import RecursiveCharacterTextSplitter
    ext = Path(file_path).suffix.lower()
    source = Path(file_path).name

    # Pick loader
    if ext == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader
        loader = PyPDFLoader(file_path)
    elif ext == ".docx":
        from langchain_community.document_loaders import Docx2txtLoader
        loader = Docx2txtLoader(file_path)
    elif ext in (".html", ".htm"):
        from langchain_community.document_loaders import BSHTMLLoader
        loader = BSHTMLLoader(file_path)
    elif ext == ".json":
        from langchain_community.document_loaders import JSONLoader
        loader = JSONLoader(file_path, jq_schema=".", text_content=False)
    else:
        # .txt, .md, .rst, .log and unknown types
        from langchain_community.document_loaders import TextLoader
        loader = TextLoader(file_path, encoding="utf-8")

    docs = loader.load()

    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    split_docs = splitter.split_documents(docs)

    return [
        {
            "content": doc.page_content,
            "source": source,
            "metadata": {**doc.metadata, "file": source},
        }
        for doc in split_docs
        if doc.page_content.strip()
    ]


# ── Public entry point ────────────────────────────────────────────────────────

def ingest_file(file_path: str) -> tuple[int, str]:
    """
    Auto-detect format, parse, chunk, embed, and store in pgvector.
    Returns (chunk_count, format_name).
    """
    from rag.store import add_chunks

    ext = Path(file_path).suffix.lower()

    if ext in (".xlsx", ".xls", ".xlsm", ".xlsb"):
        chunks = _load_excel(file_path)
        fmt = "excel"
    elif ext == ".csv":
        chunks = _load_csv(file_path)
        fmt = "csv"
    elif ext in (".pptx", ".ppt"):
        chunks = _load_pptx(file_path)
        fmt = "powerpoint"
    else:
        # PDF, Word, HTML, JSON, text — all via LangChain
        chunks = _lc_load(file_path)
        fmt = ext.lstrip(".") or "text"

    count = add_chunks(chunks)
    return count, fmt
